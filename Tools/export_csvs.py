# -*- coding: utf-8 -*-
"""
Generate categories.csv and types.csv from catalog.xlsx with UTF-8 BOM.
"""
import openpyxl
import csv

XLSX_PATH = r"D:\マイドライブ\WaveTagger\config\catalog.xlsx"
CAT_CSV_PATH = r"D:\マイドライブ\WaveTagger\config\categories.csv"
TYPE_CSV_PATH = r"D:\マイドライブ\WaveTagger\config\types.csv"

wb = openpyxl.load_workbook(XLSX_PATH)

# categories.csv
ws_cat = wb["Categories"]
with open(CAT_CSV_PATH, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    for row in ws_cat.iter_rows(values_only=True):
        if not any(row):
            continue
        # replace None with empty string
        cleaned_row = [str(c) if c is not None else "" for c in row]
        writer.writerow(cleaned_row)

print("Exported categories.csv")

# types.csv
ws_types = wb["Types"]
with open(TYPE_CSV_PATH, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    for row in ws_types.iter_rows(values_only=True):
        if not any(row):
            continue
        cleaned_row = [str(c) if c is not None else "" for c in row]
        writer.writerow(cleaned_row)

print("Exported types.csv")
