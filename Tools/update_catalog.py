# -*- coding: utf-8 -*-
"""
Script to update catalog.xlsx with extensive Categories, Matrix Types, and Types sheet Japanese translations.
"""

import openpyxl
import os
import sys
import shutil
from datetime import datetime
from build_catalog import CATALOG_DATA

EXCEL_PATH = r"D:\マイドライブ\WaveTagger\config\catalog.xlsx"

def update_catalog():
    print(f"Loading workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # 1. Update Categories Sheet
    ws_cat = wb["Categories"]
    
    # Read existing Category rows to preserve any manual notes or prefixes
    existing_cat_info = {}
    for r in list(ws_cat.iter_rows(values_only=True))[1:]:
        if not r or not r[0]:
            continue
        cat_name = str(r[0]).strip()
        ja = str(r[1]).strip() if r[1] is not None else ""
        note = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
        prefix = str(r[3]).strip() if len(r) > 3 and r[3] is not None else ""
        existing_cat_info[cat_name.upper()] = {
            "name": cat_name,
            "ja": ja,
            "note": note,
            "prefix": prefix,
        }

    # Clear Categories sheet and rewrite with new complete matrix
    ws_cat.delete_rows(1, ws_cat.max_row)
    
    header = ["Category", "Japanese", "Note", "Prefix"] + [f"TYPE{i:02d}" for i in range(1, 11)]
    ws_cat.append(header)

    # Sort categories alphabetically
    all_categories = sorted(CATALOG_DATA.keys())

    for cat in all_categories:
        def_ja, def_prefix, def_note, types = CATALOG_DATA[cat]
        
        # Preserve user custom note / prefix if existed
        if cat in existing_cat_info:
            cur = existing_cat_info[cat]
            ja = cur["ja"] if cur["ja"] else def_ja
            prefix = cur["prefix"] if cur["prefix"] else def_prefix
            note = cur["note"] if cur["note"] else def_note
        else:
            ja = def_ja
            prefix = def_prefix
            note = def_note

        row = [cat, ja, note, prefix]
        # Append types up to 10
        type_names = [t[0] for t in types][:10]
        while len(type_names) < 10:
            type_names.append("")
        row.extend(type_names)

        ws_cat.append(row)

    print(f"Categories sheet updated with {len(all_categories)} categories.")

    # 2. Update Types Sheet
    ws_types = wb["Types"]
    existing_type_rows = []
    seen_type_keys = set()

    for r in list(ws_types.iter_rows(values_only=True))[1:]:
        if not r or not r[0] or not r[1]:
            continue
        cat = str(r[0]).strip()
        typ = str(r[1]).strip()
        ja = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
        note = str(r[3]).strip() if len(r) > 3 and r[3] is not None else ""

        key = (cat.upper(), typ.upper())
        seen_type_keys.add(key)
        existing_type_rows.append([cat, typ, ja, note])

    print(f"Existing Types rows: {len(existing_type_rows)}")

    # Add missing types from CATALOG_DATA
    added_type_count = 0
    for cat in all_categories:
        _, _, _, types = CATALOG_DATA[cat]
        for typ, ja in types:
            key = (cat.upper(), typ.upper())
            if key not in seen_type_keys:
                existing_type_rows.append([cat, typ, ja, ""])
                seen_type_keys.add(key)
                added_type_count += 1
            else:
                # If Japanese translation was empty, fill it
                for item in existing_type_rows:
                    if item[0].upper() == cat.upper() and item[1].upper() == typ.upper():
                        if not item[2]:
                            item[2] = ja
                            print(f"Filled empty Japanese for {cat}/{typ}: {ja}")

    # Rewrite Types sheet
    ws_types.delete_rows(1, ws_types.max_row)
    ws_types.append(["Category", "Type", "Japanese", "Note"])
    for r in existing_type_rows:
        ws_types.append(r)

    print(f"Types sheet updated. Added {added_type_count} new type rows. Total: {len(existing_type_rows)}")

    # Save workbook
    wb.save(EXCEL_PATH)
    print("catalog.xlsx successfully saved!")

if __name__ == "__main__":
    update_catalog()
